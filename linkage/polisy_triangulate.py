# -*- coding: utf-8 -*-
"""polisy_triangulate: a triangulated, audited link from VRscores employers to Compustat firms (GVKEY).

Why
  DIPI is keyed by GVKEY, so it joins Compustat exactly. The one uncertain step in the DIPI + VRscores design is
  VRscores employer -> Compustat firm. This module turns that step from "we did fuzzy matching" into a documented
  linkage with a measured error rate, following the logic of Akcigit & Goldschlag (2025, Journal of Economic Growth):
  a link is trusted when independent sources agree, every link is labelled by how it was confirmed, selection into the
  linked sample is shown, and precision is measured by a blind hand audit.

The evidence for a pair (VRscores employer, GVKEY)
  E1  Compustat name   the employer's name matches the firm's Compustat name (POLISY module 04, high or medium tier), and
                       the name is unique on both sides: no other Compustat firm has it, and the employer holds at least
                       90% of the workers of all VRscores employers with the same cleaned name
  E2  patent route     the employer's name equals a patent owner's name (PatentsView assignee), and DISCERN 2.0 gives at
                       least 80% of that owner's patents granted 2008-2021 to this GVKEY (and links at least half of them)
  E3  DISCERN names    the employer's name is on DISCERN's own list of this firm's names and subsidiaries, and on no
                       other firm's
  E4  size             VRscores' matched workers fit the firm's Compustat employees over the years: median ratio between
                       0.2% and 150% (module 04's rule), no jump of more than 3x from one year to the next (a merger or
                       split), and the two series do not move in opposite directions
  The independent sources are Compustat names (E1), DISCERN (E2 or E3: both come from DISCERN, so they count once) and
  sizes (E4).

Match types (exactly one per pair, in this order)
  rejected               rejected by hand
  superseded             another firm was accepted by hand for this employer
  conflict               the employer reaches different firms through different sources (needs a hand decision)
  triangulated_patents   E1 and E2: two independent routes reach the same firm
  triangulated_names     E1 and E3
  name_and_size          E1 and E4 (firms whose names have no patents in 2008-2021)
  discern_and_size       E2 or E3, and E4, without E1 (subsidiaries and renamed firms, e.g. Google LLC -> Alphabet)
  hand_accepted          accepted by hand without a two-source confirmation
  single_source          one source only

Link sets (one firm per employer in every set)
  strict    triangulated_patents
  main      two independent sources: triangulated_patents, triangulated_names, name_and_size, discern_and_size
  broad     main + hand_accepted + single-source pairs in module 04's high tier
  baseline  module 04's rule as used by module 06 today (high tier, or accept = 1), for comparison

The rule that matters most: no political variable is read while links are built. build_links() reads names, years,
worker counts, Compustat names and sizes, patents and DISCERN, nothing else, and refuses a VRscores column that looks
political. Party data are read only after the links are frozen (selection_table, workforce_panels), and those steps
stop when links_all.csv no longer matches the hash recorded when it was frozen.

How to run (Colab, after the lab notebook's cells 1-3, with RUN_POLISY_DA = True, and cell 6b once)
  import polisy_triangulate as tri
  tri.build_links()          # step A: the links, the match-type tables, the review list; freezes the protocol
  tri.draw_audit()           # step B: two coder workbooks for the blind hand audit
  tri.audit_summary()        # step C: precision by match type and by link set, inter-coder agreement
  tri.selection_table()      # step D: linked versus unlinked firms (reads DIPI, so only after the freeze)
  tri.workforce_panels()     # step E: the firm-year workforce panel for each link set, for the sensitivity models
Every output goes to POLISY/linkage in Google Drive; README_outputs.txt there says what each file is.
"""
from __future__ import annotations

import hashlib
import json
import math
import os
import re
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd

try:
    import duckdb
except ImportError as e:  # pragma: no cover
    raise SystemExit("polisy_triangulate needs duckdb: run  pip install duckdb  first") from e

__version__ = "1.0 (2026-10-14)"

# ----------------------------------------------------------------------------------------------- settings
SETTINGS = {
    "match_year": 2015,             # the Compustat fiscal year of module 04's name match
    "min_workers": 25,              # VRscores employers below this in every year are left out (module 04's rule)
    "vr_dominance": 0.90,           # share of same-name workers an employer must hold for its name to count as unique
    "patent_years": (2008, 2021),   # grant years whose patents decide who owns an assignee (DISCERN ends in 2021)
    "owner_share": 0.80,            # share of an assignee's linked patents one firm must own (E2)
    "linked_share": 0.50,           # share of an assignee's patents DISCERN must link to some listed firm (E2)
    "min_patents": 1,               # linked patents an assignee needs for E2
    "flag_patents": 5,              # an assignee with this many patents, almost none owned by a listed firm, is flagged
    "size_ratio": (0.002, 1.5),     # VRscores matched workers / Compustat employees (module 04's rule)
    "size_max_jump": 3.0,           # largest year-to-year change of that ratio before a link looks like a merger or split
    "size_years": (2012, 2024),     # years compared for the size trajectory
    "window": (2012, 2021),         # years the selection table averages over (the VRscores-DISCERN overlap)
    "audit_n": 200,                 # links in the hand audit
    "audit_min_per_type": 20,       # at least this many per match type (all of them when a type has fewer)
    "seed": 20261014,               # fixes the audit sample, so it can be drawn again identically
}

TYPES = ["triangulated_patents", "triangulated_names", "name_and_size", "discern_and_size", "hand_accepted",
         "single_source", "conflict", "superseded", "rejected"]
TYPE_LABELS = {
    "triangulated_patents": "Compustat name and patent route (DISCERN) reach the same firm",
    "triangulated_names": "Compustat name and DISCERN's name list reach the same firm",
    "name_and_size": "Compustat name and size trajectory agree (no patent route for the name)",
    "discern_and_size": "DISCERN (patents or names) and size agree, without a Compustat name match",
    "hand_accepted": "accepted by hand, no two-source confirmation",
    "single_source": "one source only",
    "conflict": "sources reach different firms: needs a hand decision",
    "superseded": "another firm was accepted by hand for this employer",
    "rejected": "rejected by hand",
}
MAIN_TYPES = ["triangulated_patents", "triangulated_names", "name_and_size", "discern_and_size"]
SETS = {
    "strict": "Compustat name and patent route agree",
    "main": "two independent sources agree",
    "broad": "main + hand-accepted + one-source links of module 04's high tier",
    "baseline": "module 04's rule as used by module 06 today",
}
# Paths you set once per session, e.g. tri.PATHS["compustat"] = "/content/drive/MyDrive/POLISY/data/my_compustat.csv"
PATHS = {}
POLITICAL = re.compile(r"dem|rep|party|liberal|conserv|ideolog|share|vote|partisan")
TIER_RANK = {"high": 0, "medium": 1, "low": 2}


# ----------------------------------------------------------------------------------------------- helpers
def log(msg):
    print(f"[{time.strftime('%H:%M:%S')}] {msg}", flush=True)


def _squash(s):
    return re.sub(r"[^a-z0-9]", "", str(s).lower())


def _sql(p):
    return str(Path(p).as_posix()).replace("'", "''")


def _gvkey(s):
    """GVKEY as Compustat writes it: six digits with leading zeros (1004, '1004.0', '001004' -> '001004')."""
    s = pd.Series(s, dtype="object").astype(str).str.strip().str.replace(r"\.0+$", "", regex=True)
    ok = s.str.fullmatch(r"\d{1,6}").fillna(False).astype(bool)
    return s.where(ok).str.zfill(6)


def _pc():
    """polisy_core when the notebook has loaded it (the same name cleaning and file finder as module 04), else None."""
    mod = sys.modules.get("polisy_core")
    if mod is not None:
        return mod
    try:
        import polisy_core as mod  # noqa: F401
        return mod
    except ImportError:
        return None


# The cleaning rule of polisy_core.norm_name, used only when polisy_core is not loaded. Both sides of every comparison
# here are cleaned in this run with the same function, so the two never mix.
_LEGAL_TAIL = {"inc", "incorporated", "corp", "corporation", "co", "company", "companies", "llc",
               "ltd", "limited", "plc", "lp", "llp", "holdings", "holding", "group", "the", "and"}
_ABBR = {"international": "intl", "technologies": "tech", "technology": "tech", "systems": "sys",
         "services": "svcs", "communications": "comm", "manufacturing": "mfg",
         "pharmaceuticals": "pharma", "laboratories": "labs", "industries": "inds",
         "financial": "finl", "national": "natl", "american": "amer", "america": "amer",
         "corporation": "corp", "company": "co", "incorporated": "inc"}


def _norm_name_fallback(s):
    if not isinstance(s, str):
        return ""
    s = re.sub(r"\(.*?\)", " ", s.lower())
    s = re.sub(r"/[a-z]{2,4}/?", " ", s)
    s = re.sub(r"-\s*(cl|class)\s*[a-z]\b", " ", s)
    s = s.replace("&", " and ")
    s = re.sub(r"[^a-z0-9 ]", " ", s)
    toks = [_ABBR.get(t, t) for t in s.split()]
    out, run = [], []
    for t in toks:
        if len(t) == 1 and t.isalpha():
            run.append(t)
            continue
        if run:
            out.append("".join(run))
            run = []
        out.append(t)
    if run:
        out.append("".join(run))
    while out and out[0] == "the":
        out = out[1:]
    while out and out[-1] in _LEGAL_TAIL:
        out.pop()
    return " ".join(out)


def _norm_fn():
    pc = _pc()
    if pc is not None and hasattr(pc, "norm_name"):
        return pc.norm_name, "polisy_core.norm_name"
    return _norm_name_fallback, "built-in copy of polisy_core.norm_name"


def _clean(series, fn):
    """Clean many names quickly: each distinct name once."""
    s = pd.Series(series, dtype="object")
    uniq = pd.Series(s.dropna().unique())
    mapping = dict(zip(uniq, uniq.map(lambda x: fn(x) if isinstance(x, str) else "")))
    return s.map(mapping).fillna("")


def _settings(settings=None):
    S = dict(SETTINGS)
    S.update(settings or {})
    return S


# ----------------------------------------------------------------------------------------------- where the inputs are
def default_paths(base=None):
    """Where each input is, in the folders the lab notebook uses. Lists are tried in turn; the newest existing file wins.

    base   your POLISY folder in Google Drive (POLISY_BASE, default /content/drive/MyDrive/POLISY)
    POLISY_DA's outputs are in POLISY_ROOT (/content/polisy) during a session and saved to POLISY/polisy_da in Drive;
    the lab's are in POLISY/lab (POLISY_LAB_ROOT)."""
    base = Path(base or os.environ.get("POLISY_BASE", "/content/drive/MyDrive/POLISY"))
    root = Path(os.environ.get("POLISY_ROOT", "/content/polisy"))
    lab = Path(os.environ.get("POLISY_LAB_ROOT", base / "lab"))
    saved = base / "polisy_da"
    return {
        "base": base,
        "out": base / "linkage",
        "vr_employer": [root / "data/canonical/vr_employer.parquet", saved / "data/canonical/vr_employer.parquet"],
        "name_match": [root / "data/keys/cw_vrid_gvkey.csv", saved / "data/keys/cw_vrid_gvkey.csv"],
        "compustat": [],   # empty: found by polisy_core as in module 04, else searched for in POLISY/data
        "patents": [lab / "canonical/patents.parquet"],
        "assignees": [lab / "staged/patentsview/g_assignee_disambiguated.parquet"],
        "discern_patents": [lab / "canonical/discern_patents.parquet"],
        "discern_names": [lab / "staged/discern/names.parquet"],
        "dipi": [root / "data/canonical/dipi.parquet", saved / "data/canonical/dipi.parquet"],
        "firm_patents": [lab / "panels/firm_patents_year.parquet"],
        "hand_decisions": None,   # None: every hand_decisions*.xlsx / *.csv in the linkage folder, oldest first
    }


def _paths(base=None, paths=None):
    P = default_paths(base)
    for k, v in {**PATHS, **(paths or {})}.items():
        P[k] = v
    P["out"] = Path(P["out"])
    P["out"].mkdir(parents=True, exist_ok=True)
    return P


def _resolve(P, key, required=False, hint=""):
    v = P.get(key)
    cands = [Path(x) for x in (v if isinstance(v, (list, tuple)) else [v]) if x]
    have = [p for p in cands if p.exists()]
    if not have:
        if required:
            raise FileNotFoundError(f"{key}: not found (looked in {', '.join(str(p) for p in cands) or 'nothing'}). {hint}")
        return None
    best = max(have, key=lambda p: p.stat().st_mtime)
    return best


def _compustat_source(P):
    """(path, member): CONFIG path or search as in module 04, else a file named like Compustat in POLISY/data."""
    p = _resolve(P, "compustat")
    if p is not None:
        return p, None
    pc = _pc()
    if pc is not None:
        try:
            loc = pc.locate("COMPUSTAT")
            if loc.get("path") is not None:
                return Path(loc["path"]), loc.get("member")
        except Exception as e:  # noqa: BLE001
            log(f"compustat: polisy_core could not locate it ({type(e).__name__}: {e}); searching POLISY/data")
    data = Path(P["base"]) / "data"
    hits = []
    if data.exists():
        for f in list(data.glob("*")) + list(data.glob("*/*")):
            if f.is_file() and re.search(r"compustat|funda|comp_final", f.name, re.I) and \
                    f.suffix.lower() in (".csv", ".dta", ".parquet", ".zip", ".txt", ".tsv"):
                hits.append(f)
    if not hits:
        raise FileNotFoundError("compustat: no Compustat file found. Put it in POLISY/data, or pass "
                                "paths={'compustat': '/content/drive/MyDrive/POLISY/data/<your file>'}")
    return max(hits, key=lambda f: f.stat().st_mtime), None


def _read_compustat(src, want=("gvkey", "fyear", "conm", "state", "emp", "sale", "at", "sich")):
    path, member = src
    pc = _pc()
    want_sq = {_squash(w) for w in want}
    if pc is not None:
        df = pc.read_table(path, member, header_hint="gvkey", columns=want_sq)
    else:
        suf = Path(path).suffix.lower()
        if suf == ".parquet":
            df = pd.read_parquet(path)
        elif suf == ".dta":
            df = pd.read_stata(path, convert_categoricals=False)
        else:
            df = pd.read_csv(path, dtype=str, usecols=lambda c: _squash(c) in want_sq, low_memory=False)
        df = df[[c for c in df.columns if _squash(c) in want_sq]]
    df = df.rename(columns={c: _squash(c) for c in df.columns})
    df = df.loc[:, ~df.columns.duplicated()]
    for k in want:
        if k not in df.columns:
            df[k] = np.nan
    df["gvkey"] = _gvkey(df["gvkey"])
    for k in ("fyear", "emp", "sale", "at", "sich"):
        df[k] = pd.to_numeric(df[k], errors="coerce")
    return df.dropna(subset=["gvkey", "fyear"]).astype({"fyear": int})


def _fingerprint(p):
    if p is None:
        return None
    p = Path(p)
    st = p.stat()
    return {"path": str(p), "bytes": st.st_size,
            "modified": datetime.fromtimestamp(st.st_mtime, timezone.utc).isoformat(timespec="seconds")}


def _con(out):
    c = duckdb.connect()
    tmp = Path(out) / "_duckdb_tmp"
    tmp.mkdir(parents=True, exist_ok=True)
    c.execute(f"SET temp_directory = '{_sql(tmp)}'")
    c.execute("SET preserve_insertion_order = false")
    return c


def _pid_sql(expr):
    """Patent numbers compared across PatentsView and DISCERN (the lab's rule): digits without leading zeros, or the
    letters and digits of design, reissue and plant patents (D, RE, PP)."""
    v = f"upper(regexp_replace(CAST({expr} AS VARCHAR), '\\.0+$', ''))"
    return (f"CASE WHEN regexp_matches({v}, '^[0-9]+$') THEN ltrim({v}, '0') "
            f"ELSE regexp_replace({v}, '[^A-Z0-9]', '', 'g') END")


def _cols(c, path):
    return [r[0] for r in c.execute(f"DESCRIBE SELECT * FROM read_parquet('{_sql(path)}')").fetchall()]


def _pick(cols, *names):
    sq = {_squash(x): x for x in cols}
    return next((sq[n] for n in names if n in sq), None)


INPUTS = [
    ("vr_employer", "VRscores employer panel (Parquet)", "POLISY_DA module 02: lab notebook cell 3 with RUN_POLISY_DA = True", "A, E"),
    ("name_match", "module 04's name match (cw_vrid_gvkey.csv)", "POLISY_DA module 04: lab notebook cell 3", "A"),
    ("compustat", "Compustat annual file", "your WRDS download, in POLISY/data", "A, D"),
    ("patents", "patents with grant year and first owner", "lab notebook cell 6b (RUN_PATENT_LAYER = True)", "A (E2)"),
    ("assignees", "every patent owner (PatentsView)", "lab notebook cell 6b (optional: else the first owner is used)", "A (E2)"),
    ("discern_patents", "DISCERN: which Compustat firm owns each patent", "lab notebook cell 6b, from your DISCERN files", "A (E2)"),
    ("discern_names", "DISCERN's firm and subsidiary names", "lab notebook cell 6b (optional)", "A (E3)"),
    ("dipi", "DIPI firm-year file (Parquet)", "POLISY_DA module 01 or 07 (cell 3)", "D only"),
    ("firm_patents", "the lab's firm-year patent panel", "lab notebook cell 6b (optional)", "D only"),
]


def check_inputs(base=None, paths=None):
    """Which input file each step will use, and what makes the missing ones. Reads nothing but file names."""
    P = _paths(base, paths)
    rows = []
    for key, what, made_by, steps in INPUTS:
        try:
            if key == "compustat":
                f, member = _compustat_source(P)
                f = f"{f}" + (f" [{member}]" if member else "")
            else:
                f = _resolve(P, key)
        except FileNotFoundError:
            f = None
        rows.append({"input": what, "status": "found" if f else "MISSING", "file": str(f) if f else "",
                     "made by": made_by, "needed for step": steps})
    files = _decision_files(P)
    rows.append({"input": "hand decisions (hand_decisions*.xlsx)", "status": f"{len(files)} file(s)" if files else "none yet",
                 "file": ", ".join(f.name for f in files), "made by": "you, from review_list.xlsx", "needed for step": "A"})
    df = pd.DataFrame(rows)
    with pd.option_context("display.width", 220, "display.max_colwidth", 90, "display.max_columns", 10):
        print(df.to_string(index=False))
    need = df.loc[df["needed for step"].isin(["A", "A, E", "A, D"]) & df.status.eq("MISSING"), "input"].tolist()
    if need:
        log("Step A cannot run yet; missing: " + "; ".join(need))
    elif df.loc[df.input.str.startswith(("patents with", "DISCERN: which")), "status"].eq("MISSING").any():
        log("Step A can run, but without the patent route (E2): run the lab notebook's cell 6b first for the full design")
    else:
        log("All inputs for step A are in place")
    return df


# ----------------------------------------------------------------------------------------------- step A: build the links
def _vr_employers(c, path, S, norm):
    """VRscores employers: id, name, years and matched-worker counts. Party columns are never selected."""
    cols = _cols(c, path)
    m = {"vrid": _pick(cols, "vrid", "unit", "employerid"), "year": _pick(cols, "year"),
         "name": _pick(cols, "companyname", "employername", "name"),
         "workers": _pick(cols, "employeecount", "workers", "matchedworkers")}
    missing = [k for k, v in m.items() if v is None]
    if missing:
        raise ValueError(f"VRscores employer panel lacks {missing}; its columns are {cols[:20]}")
    for k, v in m.items():
        if POLITICAL.search(_squash(v)):
            raise RuntimeError(f"refusing to read VRscores column {v!r} while building links (it looks political)")
    q = lambda x: '"' + x.replace('"', '""') + '"'  # noqa: E731
    c.execute(f"""CREATE OR REPLACE TABLE vr AS
        SELECT CAST({q(m['vrid'])} AS VARCHAR) AS vrid, CAST({q(m['year'])} AS INTEGER) AS year,
               CAST({q(m['name'])} AS VARCHAR) AS name, CAST({q(m['workers'])} AS DOUBLE) AS workers
        FROM read_parquet('{_sql(path)}')""")
    # the name in the match year (the one module 04 matched on), else the latest name
    emp = c.execute(f"""SELECT vrid, coalesce(max(name) FILTER (WHERE year = {int(S['match_year'])}), arg_max(name, year))
                                   AS vr_name, max(workers) AS vr_workers_max,
                               max(workers) FILTER (WHERE year = {int(S['match_year'])}) AS vr_workers_match_year,
                               min(year) AS vr_first_year, max(year) AS vr_last_year
                        FROM vr WHERE name IS NOT NULL GROUP BY vrid""").df()
    emp["name_norm"] = _clean(emp.vr_name, norm)
    emp["vr_workers_max"] = emp.vr_workers_max.fillna(0)
    g = emp[emp.name_norm != ""].groupby("name_norm").vr_workers_max
    emp["vr_same_name_n"] = emp.name_norm.map(g.size()).fillna(0).astype(int)
    tot = emp.name_norm.map(g.sum())
    emp["vr_name_dominance"] = np.where(tot > 0, emp.vr_workers_max / tot, np.where(emp.vr_same_name_n == 1, 1.0, 0.0))
    emp["vr_name_ok"] = (emp.name_norm != "") & ((emp.vr_same_name_n == 1) | (emp.vr_name_dominance >= S["vr_dominance"]))
    log(f"VRscores: {len(emp):,} employers; {(emp.vr_workers_max >= S['min_workers']).sum():,} with at least "
        f"{S['min_workers']} matched workers in some year; {int((emp.vr_same_name_n > 1).sum()):,} share a cleaned name with "
        f"another employer ({int(((emp.vr_same_name_n > 1) & emp.vr_name_ok).sum()):,} of them hold "
        f"{S['vr_dominance']:.0%} or more of that name's workers)")
    return emp


def _route1(path):
    cw = pd.read_csv(path, dtype={"gvkey": str, "vrid": str}, low_memory=False)
    need = {"gvkey", "vrid", "tier"}
    if not need <= set(cw.columns):
        raise ValueError(f"{path.name} lacks {sorted(need - set(cw.columns))}: rebuild it with module 04")
    for k in ("method", "score", "accept"):
        if k not in cw.columns:
            cw[k] = np.nan
    cw["gvkey"] = _gvkey(cw.gvkey)
    cw["vrid"] = cw.vrid.astype(str).str.strip()
    cw["accept"] = pd.to_numeric(cw.accept, errors="coerce")
    cw["score"] = pd.to_numeric(cw.score, errors="coerce")
    cw["_rank"] = cw.tier.map(TIER_RANK).fillna(9)
    cw = cw.dropna(subset=["gvkey"]).sort_values(["vrid", "gvkey", "_rank", "score"], ascending=[True, True, True, False])
    cw = cw.drop_duplicates(["vrid", "gvkey"], keep="first")
    out = cw[["vrid", "gvkey", "tier", "method", "score", "accept"]].rename(
        columns={"tier": "r1_tier", "method": "r1_method", "score": "r1_score", "accept": "cw_accept"})
    out["r1"] = True
    log(f"route 1 (module 04's name match): {len(out):,} candidate pairs, "
        + ", ".join(f"{k} {v:,}" for k, v in out.r1_tier.value_counts().items())
        + f"; accept filled by hand for {int(out.cw_accept.notna().sum()):,}")
    return out


def _decision_files(P):
    v = P.get("hand_decisions")
    if v:
        files = [Path(x) for x in (v if isinstance(v, (list, tuple)) else [v]) if Path(x).exists()]
    else:
        files = [f for f in P["out"].glob("hand_decisions*") if f.suffix.lower() in (".csv", ".xlsx")
                 and not f.name.startswith("~$")]
    return sorted(files, key=lambda f: (f.stat().st_mtime, f.name))


def _read_decisions(f):
    if f.suffix.lower() == ".xlsx":
        book = pd.ExcelFile(f)
        sheet = "Review" if "Review" in book.sheet_names else book.sheet_names[0]
        d = book.parse(sheet, dtype=str)
    else:
        d = pd.read_csv(f, dtype=str)
    d.columns = [str(c).strip().lower() for c in d.columns]
    if not {"vrid", "gvkey", "accept"} <= set(d.columns):
        raise ValueError(f"{f.name} needs the columns vrid, gvkey and accept (1 or 0)")
    d["gvkey"] = _gvkey(d.gvkey)
    d["vrid"] = d.vrid.astype(str).str.strip()
    d["accept"] = pd.to_numeric(d.accept.astype(str).str.strip().str.replace(r"\.0+$", "", regex=True), errors="coerce")
    return d[d.accept.isin([0, 1])].dropna(subset=["gvkey"])[["vrid", "gvkey", "accept"]]


def _hand(P, r1):
    """Hand decisions: accept in module 04's file, overridden by every hand_decisions*.xlsx / *.csv in the linkage
    folder (columns vrid, gvkey, accept), oldest first, so a later round overrides an earlier one."""
    h = r1.loc[r1.cw_accept.isin([0, 1]), ["vrid", "gvkey", "cw_accept"]].rename(columns={"cw_accept": "accept"})
    h["hand_source"] = "name-match file"
    for f in _decision_files(P):
        d = _read_decisions(f).drop_duplicates(["vrid", "gvkey"], keep="last")
        d["hand_source"] = f.name
        h = pd.concat([h, d]).drop_duplicates(["vrid", "gvkey"], keep="last")
        log(f"hand decisions: {len(d):,} from {f.name}")
    return h.rename(columns={"accept": "hand"})


def _patent_route(c, P, S, norm, vr_names):
    """Each cleaned assignee name: its patents granted in the window, how many DISCERN gives to a listed firm, and the
    firm owning most of them. Only names that some VRscores employer also has are kept."""
    pats = _resolve(P, "patents")
    disc = _resolve(P, "discern_patents")
    if pats is None or disc is None:
        log("patent route: off. It needs the lab's patent layer (cell 6b): "
            + ", ".join(f"{k} {'found' if v else 'MISSING'}" for k, v in (("lab/canonical/patents.parquet", pats),
                                                                        ("lab/canonical/discern_patents.parquet", disc))))
        return None, {}
    y0, y1 = S["patent_years"]
    pcols = _cols(c, pats)
    if "year" not in pcols or "patent_id" not in pcols:
        raise ValueError(f"{pats.name} lacks patent_id or year; rebuild the lab's patent layer")
    c.execute(f"""CREATE OR REPLACE TABLE pw AS SELECT DISTINCT CAST(patent_id AS VARCHAR) AS patent_id,
                  {_pid_sql('patent_id')} AS pid FROM read_parquet('{_sql(pats)}') WHERE year BETWEEN {y0} AND {y1}""")
    asg = _resolve(P, "assignees")
    if asg is not None:
        acols = _cols(c, asg)
        org, at = _pick(acols, "org", "disambigassigneeorganization", "organization"), _pick(acols, "atype", "assigneetype")
        if org is None:
            asg = None
    if asg is not None:
        typ = f"(\"{at}\" IS NULL OR try_cast(try_cast(\"{at}\" AS DOUBLE) AS INTEGER) % 10 IN (2, 3))" if at else "TRUE"
        c.execute(f"""CREATE OR REPLACE TABLE ao AS SELECT DISTINCT CAST(patent_id AS VARCHAR) AS patent_id,
                      CAST("{org}" AS VARCHAR) AS org FROM read_parquet('{_sql(asg)}')
                      WHERE "{org}" IS NOT NULL AND {typ}""")
        src = "every assignee (PatentsView g_assignee_disambiguated)"
    else:
        if "assignee_org" not in pcols:
            log("patent route: off (no assignee names in the lab's patents table)")
            return None, {}
        typ = "(assignee_type IS NULL OR assignee_type IN (2, 3))" if "assignee_type" in pcols else "TRUE"
        c.execute(f"""CREATE OR REPLACE TABLE ao AS SELECT DISTINCT CAST(patent_id AS VARCHAR) AS patent_id,
                      CAST(assignee_org AS VARCHAR) AS org FROM read_parquet('{_sql(pats)}')
                      WHERE assignee_org IS NOT NULL AND {typ}""")
        src = "the first assignee of each patent (the lab's patents table)"
    dcols = _cols(c, disc)
    if "gvkey" not in dcols or "patent_id" not in dcols:
        raise ValueError(f"{disc.name} lacks patent_id or gvkey; rebuild the lab's patent layer")
    c.execute(f"""CREATE OR REPLACE TABLE dp AS SELECT DISTINCT {_pid_sql('patent_id')} AS pid,
                  lpad(regexp_extract(CAST(gvkey AS VARCHAR), '(\\d+)', 1), 6, '0') AS gvkey
                  FROM read_parquet('{_sql(disc)}') WHERE gvkey IS NOT NULL""")
    orgs = c.execute("SELECT a.org, count(*) AS n FROM ao a JOIN pw USING (patent_id) GROUP BY 1").df()
    orgs["name_norm"] = _clean(orgs.org, norm)
    orgs = orgs[(orgs.name_norm != "") & orgs.name_norm.isin(vr_names)]
    example = (orgs.sort_values(["name_norm", "n", "org"], ascending=[True, False, True])
               .drop_duplicates("name_norm").set_index("name_norm").org)
    c.register("orgmap_df", orgs[["org", "name_norm"]])
    c.execute("CREATE OR REPLACE TABLE orgmap AS SELECT * FROM orgmap_df")
    c.execute("""CREATE OR REPLACE TABLE an AS SELECT DISTINCT m.name_norm, w.pid
                 FROM ao a JOIN pw w USING (patent_id) JOIN orgmap m USING (org)""")
    c.execute("CREATE OR REPLACE TABLE ang AS SELECT an.name_norm, an.pid, d.gvkey FROM an LEFT JOIN dp d USING (pid)")
    tot = c.execute("""SELECT name_norm, count(DISTINCT pid) AS pr_total,
                              count(DISTINCT pid) FILTER (WHERE gvkey IS NOT NULL) AS pr_linked
                       FROM ang GROUP BY 1""").df()
    own = c.execute("SELECT name_norm, gvkey, count(DISTINCT pid) AS n FROM ang WHERE gvkey IS NOT NULL GROUP BY 1, 2").df()
    own = own.sort_values(["name_norm", "n", "gvkey"], ascending=[True, False, True])
    n_owners = own.groupby("name_norm").size().rename("pr_owners")
    dom = own.drop_duplicates("name_norm").rename(columns={"gvkey": "pr_owner", "n": "pr_owner_patents"})
    r = tot.merge(dom, on="name_norm", how="left").merge(n_owners, on="name_norm", how="left")
    r["pr_owners"] = r.pr_owners.fillna(0).astype(int)
    r["pr_owner_share"] = r.pr_owner_patents / r.pr_linked.where(r.pr_linked > 0)
    r["pr_linked_share"] = r.pr_linked / r.pr_total.where(r.pr_total > 0)
    r["pr_example_assignee"] = r.name_norm.map(example)
    r["pr_usable"] = ((r.pr_linked >= S["min_patents"]) & (r.pr_owner_share >= S["owner_share"])
                      & (r.pr_linked_share >= S["linked_share"]))
    log(f"patent route: {src}; {len(r):,} cleaned assignee names shared with VRscores employers, "
        f"{int(r.pr_usable.sum()):,} with one dominant DISCERN owner (patents granted {y0}-{y1})")
    return r.set_index("name_norm"), {"patents": _fingerprint(pats), "discern_patents": _fingerprint(disc),
                                      "assignees": _fingerprint(asg)}


def _discern_names(P, S, norm, vr_names):
    f = _resolve(P, "discern_names")
    if f is None:
        log("DISCERN name list: not found (the lab stages it in cell 6b); E3 is off")
        return None, {}
    d = pd.read_parquet(f)
    gv, nm = _pick(d.columns, "gvkey"), _pick(d.columns, "name")
    if gv is None or nm is None:
        log(f"DISCERN name list {f.name}: no gvkey or name column; E3 is off")
        return None, {}
    d = d.rename(columns={gv: "gvkey", nm: "name"})
    d["gvkey"] = _gvkey(d.gvkey)
    if "fy1" in d.columns:
        fy1 = pd.to_numeric(d.fy1, errors="coerce")
        d = d[fy1.isna() | (fy1 >= S["patent_years"][0])]
    d["name_norm"] = _clean(d.name, norm)
    d = d[(d.name_norm != "") & d.name_norm.isin(vr_names)].dropna(subset=["gvkey"])
    g = d.groupby("name_norm").gvkey.agg(lambda s: sorted(set(s)))
    out = pd.DataFrame({"dn_firms": g.map(len), "dn_gvkey": g.map(lambda v: v[0] if len(v) == 1 else None)})
    log(f"DISCERN name list: {len(out):,} cleaned names shared with VRscores employers, "
        f"{int((out.dn_firms == 1).sum()):,} on one firm's list only")
    return out, {"discern_names": _fingerprint(f)}


def _size_evidence(c, pairs, comp, S):
    """Size trajectory of each candidate pair: VRscores matched workers against Compustat employees, year by year."""
    y0, y1 = S["size_years"]
    c.register("pairs_df", pairs[["vrid"]].drop_duplicates())
    vy = c.execute(f"""SELECT v.vrid, v.year, max(v.workers) AS workers FROM vr v JOIN pairs_df p USING (vrid)
                       WHERE v.year BETWEEN {y0} AND {y1} GROUP BY 1, 2""").df()
    cy = (comp[(comp.fyear >= y0) & (comp.fyear <= y1)].dropna(subset=["emp"])
          .sort_values(["gvkey", "fyear", "emp"]).drop_duplicates(["gvkey", "fyear"], keep="last")
          .rename(columns={"fyear": "year"})[["gvkey", "year", "emp"]])
    m = pairs[["vrid", "gvkey"]].drop_duplicates().merge(vy, on="vrid").merge(cy, on=["gvkey", "year"])
    m = m[(m.workers > 0) & (m.emp > 0)].copy()
    res = pd.DataFrame(columns=["vrid", "gvkey", "size_years", "size_ratio_median", "size_max_jump", "size_corr"])
    if len(m):
        m["ratio"] = m["workers"] / (m["emp"] * 1000.0)
        m["log_r"], m["log_w"], m["log_e"] = np.log(m["ratio"]), np.log(m["workers"]), np.log(m["emp"])
        m = m.sort_values(["vrid", "gvkey", "year"])
        m["d_log_r"] = m.groupby(["vrid", "gvkey"], sort=False)["log_r"].diff().abs()
        m["log_we"], m["log_w2"], m["log_e2"] = m["log_w"] * m["log_e"], m["log_w"] ** 2, m["log_e"] ** 2
        a = m.groupby(["vrid", "gvkey"]).agg(size_years=("year", "size"), size_ratio_median=("ratio", "median"),
                                              jump=("d_log_r", "max"), mw=("log_w", "mean"), me=("log_e", "mean"),
                                              mwe=("log_we", "mean"), mw2=("log_w2", "mean"), me2=("log_e2", "mean"))
        vw, ve = a["mw2"] - a["mw"] ** 2, a["me2"] - a["me"] ** 2
        cov = a["mwe"] - a["mw"] * a["me"]
        ok = (a["size_years"] >= 4) & (vw > 1e-9) & (ve > 1e-9)
        a["size_corr"] = np.where(ok, cov / np.sqrt(vw.clip(lower=1e-12) * ve.clip(lower=1e-12)), np.nan)
        a["size_max_jump"] = np.exp(a["jump"])
        res = a.reset_index()[["vrid", "gvkey", "size_years", "size_ratio_median", "size_max_jump", "size_corr"]]
    out = pairs.merge(res, on=["vrid", "gvkey"], how="left")
    for k in ("size_years", "size_ratio_median", "size_max_jump", "size_corr"):
        out[k] = pd.to_numeric(out[k], errors="coerce")
    out["size_years"] = out.size_years.fillna(0).astype(int)
    lo, hi = S["size_ratio"]
    in_range = out.size_ratio_median.between(lo, hi)
    no_break = out.size_max_jump.isna() | (out.size_max_jump <= S["size_max_jump"])
    not_opposite = out.size_corr.isna() | (out.size_corr >= 0)
    out["e4"] = (out.size_years >= 1) & in_range & no_break & not_opposite
    out["size_status"] = np.select(
        [out.size_years == 0, ~in_range, ~no_break, ~not_opposite],
        ["no overlapping years", "ratio out of range", "break (merger or split?)", "moves in the opposite direction"],
        "consistent")
    return out


def _one_firm_per_employer(df, col):
    """In each link set an employer belongs to one firm: when an employer reaches two, keep the best-evidenced pair."""
    rank = {t: i for i, t in enumerate(TYPES)}
    s = df[df[col]].assign(_r=lambda x: x.match_type.map(rank), _w=lambda x: -x.vr_workers_max.fillna(0))
    s = s.sort_values(["vrid", "_r", "_w", "gvkey"])
    dup = s.duplicated("vrid", keep="first")
    if dup.any():
        drop = s.index[dup]
        df.loc[drop, col] = False
        log(f"set {col[4:]}: {len(drop):,} pairs dropped because their employer has a better-evidenced pair")
    return df


def build_links(base=None, paths=None, settings=None, note=""):
    """Step A. Build every candidate pair, its evidence and match type, the link sets, the tables and the review list,
    and freeze the result (protocol.json). Reads no political variable."""
    t0 = time.time()
    S = _settings(settings)
    P = _paths(base, paths)
    out = P["out"]
    norm, norm_src = _norm_fn()
    log(f"polisy_triangulate {__version__}: building links into {out} (names cleaned with {norm_src})")
    vr_path = _resolve(P, "vr_employer", True, "Run the lab notebook's cell 3 with RUN_POLISY_DA = True "
                                               "(module 02 builds it), so POLISY/polisy_da holds the VRscores Parquet.")
    cw_path = _resolve(P, "name_match", True, "Module 04 builds keys/cw_vrid_gvkey.csv (cell 3, RUN_POLISY_DA = True).")
    comp_src = _compustat_source(P)
    log(f"inputs: {vr_path} | {cw_path} | {comp_src[0]}" + (f" [{comp_src[1]}]" if comp_src[1] else ""))
    c = _con(out)
    try:
        emp = _vr_employers(c, vr_path, S, norm)
        comp = _read_compustat(comp_src)
        if comp.emp.notna().sum() == 0:
            log("WARNING Compustat file has no 'emp' column: the size trajectory (E4) cannot be checked")
        cy = comp[comp.fyear == S["match_year"]].sort_values(["gvkey"]).drop_duplicates("gvkey")
        cy = cy.assign(comp_name_norm=_clean(cy.conm, norm))
        nn = cy[cy.comp_name_norm != ""].groupby("comp_name_norm").gvkey.size()
        cy["comp_name_n"] = cy.comp_name_norm.map(nn).fillna(0).astype(int)
        latest = comp.dropna(subset=["conm"]).sort_values(["gvkey", "fyear"]).drop_duplicates("gvkey", keep="last")
        firm = latest[["gvkey", "conm", "state"]].rename(columns={"conm": "comp_name", "state": "comp_state"})
        firm = firm.merge(cy[["gvkey", "conm", "emp", "comp_name_n"]].rename(
            columns={"conm": "comp_name_match_year", "emp": "comp_emp_match_year"}), on="gvkey", how="left")
        log(f"Compustat: {comp.gvkey.nunique():,} firms; {len(cy):,} in fiscal {S['match_year']}, "
            f"{int((cy.comp_name_n > 1).sum()):,} of them sharing a cleaned name with another firm")

        r1 = _route1(cw_path)
        hand = _hand(P, r1)
        vr_names = set(emp.loc[emp.vr_name_ok, "name_norm"])
        pr, fp_pr = _patent_route(c, P, S, norm, vr_names)
        dn, fp_dn = _discern_names(P, S, norm, vr_names)

        pool = emp[(emp.vr_workers_max >= S["min_workers"]) & emp.vr_name_ok]
        cand = [r1[["vrid", "gvkey"]]]
        if pr is not None and len(pr):
            use = pr[pr.pr_usable]
            x = pool[pool.name_norm.isin(use.index)]
            cand.append(pd.DataFrame({"vrid": x.vrid.values, "gvkey": x.name_norm.map(use.pr_owner).values}))
        if dn is not None and len(dn):
            one = dn[dn.dn_firms == 1]
            x = pool[pool.name_norm.isin(one.index)]
            cand.append(pd.DataFrame({"vrid": x.vrid.values, "gvkey": x.name_norm.map(one.dn_gvkey).values}))
        cand.append(hand.loc[hand.hand == 1, ["vrid", "gvkey"]])
        pairs = pd.concat(cand, ignore_index=True).dropna().drop_duplicates()
        pairs["vrid"] = pairs.vrid.astype(str)

        # evidence
        df = pairs.merge(emp, on="vrid", how="left").merge(firm, on="gvkey", how="left")
        df = df.merge(r1, on=["vrid", "gvkey"], how="left").merge(hand, on=["vrid", "gvkey"], how="left")
        df["r1"] = df.r1.eq(True)
        df["vr_name_ok"] = df.vr_name_ok.eq(True)
        missing_vr = int(df.vr_name.isna().sum())
        if missing_vr:
            log(f"WARNING {missing_vr:,} candidate pairs name a VRscores employer that is not in the employer panel")
        comp_unique = df.comp_name_n.isna() | (df.comp_name_n <= 1)
        df["e1"] = df.r1 & df.r1_tier.isin(["high", "medium"]) & df.vr_name_ok & comp_unique
        if pr is not None and len(pr):
            df = df.merge(pr.drop(columns=["pr_usable"]), left_on="name_norm", right_index=True, how="left")
            usable = df.name_norm.map(pr.pr_usable).eq(True)
            df["e2"] = df.vr_name_ok & usable & df.pr_owner.eq(df.gvkey)
            df["flag_patents_unowned"] = (df.vr_name_ok & (df.pr_total >= S["flag_patents"])
                                          & (df.pr_linked_share < 0.2) & ~df.e2)
        else:
            for k in ("pr_total", "pr_linked", "pr_owner", "pr_owner_patents", "pr_owners", "pr_owner_share",
                      "pr_linked_share", "pr_example_assignee"):
                df[k] = np.nan
            df["e2"] = False
            df["flag_patents_unowned"] = False
        if dn is not None and len(dn):
            df = df.merge(dn, left_on="name_norm", right_index=True, how="left")
            df["e3"] = df.vr_name_ok & df.dn_gvkey.eq(df.gvkey)
        else:
            df["dn_firms"], df["dn_gvkey"], df["e3"] = np.nan, None, False
        df = _size_evidence(c, df, comp, S)
        df["sources"] = df.e1.astype(int) + (df.e2 | df.e3).astype(int) + df.e4.astype(int)

        # conflicts and hand decisions, employer by employer
        has_ev = df.e1 | df.e2 | df.e3 | df.hand.eq(1)
        nfirms = df[has_ev].groupby("vrid").gvkey.nunique()
        df["employer_firms"] = df.vrid.map(nfirms).fillna(0).astype(int)
        acc = df[df.hand.eq(1)].groupby("vrid").gvkey.nunique()
        df["hand_accepted_firms"] = df.vrid.map(acc).fillna(0).astype(int)
        resolved = df.hand_accepted_firms == 1
        superseded = resolved & ~df.hand.eq(1)
        conflict = (df.employer_firms > 1) & ~resolved
        df["match_type"] = np.select(
            [df.hand.eq(0), superseded, conflict & has_ev,
             df.e1 & df.e2, df.e1 & df.e3, df.e1 & df.e4, (df.e2 | df.e3) & df.e4, df.hand.eq(1)],
            ["rejected", "superseded", "conflict",
             "triangulated_patents", "triangulated_names", "name_and_size", "discern_and_size", "hand_accepted"],
            "single_source")

        # link sets: an employer in an unresolved conflict stays out of every set until a person decides
        blocked = df.vrid.isin(set(df.loc[df.match_type.eq("conflict"), "vrid"]))
        df["set_strict"] = df.match_type.eq("triangulated_patents") & ~blocked
        df["set_main"] = df.match_type.isin(MAIN_TYPES) & ~blocked
        df["set_broad"] = (df.set_main | df.match_type.eq("hand_accepted") | (
            df.match_type.eq("single_source") & df.r1_tier.eq("high"))) & ~blocked
        df["set_baseline"] = df.r1 & (df.cw_accept.eq(1) | (df.cw_accept.isna() & df.r1_tier.eq("high")))
        for col in ("set_strict", "set_main", "set_broad"):
            df = _one_firm_per_employer(df, col)

        # write
        keep = ["vrid", "gvkey", "match_type", "sources", "e1", "e2", "e3", "e4",
                "set_strict", "set_main", "set_broad", "set_baseline",
                "vr_name", "name_norm", "vr_workers_max", "vr_workers_match_year", "vr_first_year", "vr_last_year",
                "vr_same_name_n", "vr_name_dominance", "vr_name_ok",
                "comp_name", "comp_name_match_year", "comp_state", "comp_emp_match_year", "comp_name_n",
                "r1", "r1_tier", "r1_method", "r1_score", "cw_accept", "hand", "hand_source",
                "pr_example_assignee", "pr_total", "pr_linked", "pr_linked_share", "pr_owner", "pr_owner_patents",
                "pr_owner_share", "pr_owners", "flag_patents_unowned", "dn_firms", "dn_gvkey",
                "size_status", "size_years", "size_ratio_median", "size_max_jump", "size_corr",
                "employer_firms", "hand_accepted_firms"]
        names = firm.set_index("gvkey").comp_name
        df["pr_owner_name"] = df.pr_owner.map(names)
        keep.insert(keep.index("pr_owner_patents"), "pr_owner_name")
        links = df[keep].sort_values(["vrid", "gvkey"]).reset_index(drop=True)
        for col in links.select_dtypes("float").columns:
            links[col] = links[col].round(6)
        lp = out / "links_all.csv"
        links.to_csv(lp, index=False)
        for s in SETS:
            cw = links.loc[links[f"set_{s}"], ["gvkey", "vrid", "match_type"]].assign(accept=1)
            cw.to_csv(out / f"crosswalk_{s}.csv", index=False)
        _tables(links, out, S, float(emp.vr_workers_max.sum()))
        _review_list(links, out)
        proto = _freeze(out, lp, S, note, {"vr_employer": _fingerprint(vr_path), "name_match": _fingerprint(cw_path),
                                          "compustat": _fingerprint(comp_src[0]), **fp_pr, **fp_dn,
                                          "hand_decisions": [_fingerprint(f) for f in _decision_files(P)]}, norm_src)
        _readme(out)
    finally:
        c.close()
    counts = links.match_type.value_counts()
    log("match types: " + ", ".join(f"{t} {int(counts.get(t, 0)):,}" for t in TYPES if counts.get(t, 0)))
    for s in SETS:
        x = links[links[f"set_{s}"]]
        log(f"set {s:8s}: {len(x):,} employer-firm links, {x.gvkey.nunique():,} firms ({SETS[s]})")
    log(f"links frozen as version {proto['version']} (hash {proto['links_sha256'][:12]}) in {time.time() - t0:.0f}s. "
        f"Next: decide the pairs in {out / 'review_list.xlsx'} if it has rows, then draw_audit()")
    return links


# ----------------------------------------------------------------------------------------------- tables
def _sum_workers(x):
    return float(x.drop_duplicates("vrid").vr_workers_max.fillna(0).sum())


def _tables(links, out, S, total_vr):
    rows = []
    for t in TYPES:
        x = links[links.match_type == t]
        rows.append({"match_type": t, "meaning": TYPE_LABELS[t], "pairs": len(x), "employers": x.vrid.nunique(),
                     "firms": x.gvkey.nunique(), "matched_workers": _sum_workers(x),
                     "share_of_pairs": len(x) / max(len(links), 1),
                     **{f"in_{s}": int(x[f"set_{s}"].sum()) for s in SETS}})
    T1 = pd.DataFrame(rows)
    T1.to_csv(out / "T1_match_types.csv", index=False)
    srows = []
    for s in SETS:
        x = links[links[f"set_{s}"]]
        per_firm = x.groupby("gvkey").vrid.nunique()
        srows.append({"set": s, "definition": SETS[s], "links": len(x), "employers": x.vrid.nunique(),
                      "firms": x.gvkey.nunique(), "matched_workers": _sum_workers(x),
                      "share_of_vrscores_workers": _sum_workers(x) / total_vr if total_vr else np.nan,
                      "firms_with_1_employer": int((per_firm == 1).sum()),
                      "firms_with_2_employers": int((per_firm == 2).sum()),
                      "firms_with_3to5_employers": int(per_firm.between(3, 5).sum()),
                      "firms_with_6plus_employers": int((per_firm >= 6).sum()),
                      "employers_with_2plus_firms": int((x.groupby("vrid").gvkey.nunique() > 1).sum())})
    pd.DataFrame(srows).to_csv(out / "T1b_link_sets.csv", index=False)

    # two-direction agreement (Akcigit & Goldschlag's Table 3, in our setting)
    usable_pr = links.pr_owner.notna() & (links.pr_owner_share >= S["owner_share"]) & \
        (links.pr_linked_share >= S["linked_share"]) & (links.pr_linked >= S["min_patents"])
    a_rows = []
    e1 = links[links.e1 & usable_pr]
    a_rows.append({"check": "Compustat-name links (E1) whose employer name has a dominant DISCERN patent owner: "
                            "the patent route reaches the same firm", "compared": len(e1),
                   "agree": int((e1.pr_owner == e1.gvkey).sum())})
    e1v = links[links.e1].groupby("vrid").gvkey.agg(set)
    e2 = links[links.e2]
    both = e2[e2.vrid.isin(e1v.index)]
    a_rows.append({"check": "Patent-route links (E2) whose employer also has a Compustat-name link (E1): same firm",
                   "compared": len(both), "agree": int(sum(g in e1v[v] for v, g in zip(both.vrid, both.gvkey)))})
    e1d = links[links.e1 & links.dn_firms.eq(1)]
    a_rows.append({"check": "Compustat-name links (E1) whose employer name is on one firm's DISCERN name list: same firm",
                   "compared": len(e1d), "agree": int((e1d.dn_gvkey == e1d.gvkey).sum())})
    for lab_, mask in (("Compustat-name links (E1)", links.e1), ("Patent-route links (E2)", links.e2)):
        x = links[mask & (links.size_years > 0)]
        a_rows.append({"check": f"{lab_} with Compustat employees in some year: size trajectory consistent (E4)",
                       "compared": len(x), "agree": int(x.e4.sum())})
    T3 = pd.DataFrame(a_rows)
    T3["share"] = T3.agree / T3.compared.where(T3.compared > 0)
    T3.to_csv(out / "T3_agreement.csv", index=False)


REVIEW_COLUMNS = [("vrid", "vrid", 20), ("gvkey", "gvkey", 9), ("accept", "accept", 9), ("decided_by", "decided_by", 12),
                  ("note", "note", 30), ("match_type", "match type", 14), ("vr_name", "VRscores employer", 36),
                  ("vr_workers_max", "VRscores matched workers (largest year)", 14), ("vr_years", "VRscores years", 11),
                  ("comp_name", "Compustat company", 34), ("comp_state", "Compustat state", 9),
                  ("comp_emp_match_year", "Compustat employees 2015 (thousands)", 14), ("r1_tier", "name-match tier", 10),
                  ("pr_example_assignee", "patent owner with this name", 30), ("pr_owner_name", "DISCERN: firm owning its patents", 30),
                  ("pr_owner_share", "share of those patents", 10), ("size_status", "size trajectory", 22),
                  ("flag_patents_unowned", "namesake's patents have no listed owner", 12)]


def _review_list(links, out):
    """Pairs that need a person: conflicts, and one-source pairs that would enter the broad or baseline set.
    Written as an Excel file whose id columns are text, so Excel cannot strip GVKEY zeros or mangle VRscores ids."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    need = links.match_type.eq("conflict") | (links.match_type.eq("single_source") & (links.set_broad | links.set_baseline))
    r = links.loc[need].sort_values(["vr_workers_max", "vrid", "gvkey"], ascending=[False, True, True]).copy()
    r["vr_years"] = r.vr_first_year.astype("Int64").astype(str) + "-" + r.vr_last_year.astype("Int64").astype(str)
    path = out / "review_list.xlsx"
    wb = Workbook()
    hw = wb.active
    hw.title = "How to decide"
    f, fb = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    fill = PatternFill("solid", fgColor="FFF2CC")
    text = [("Linkage review list", Font(name="Arial", size=13, bold=True)), ("Prepared for the Research Team", f), ("", f),
            ("Each row is a VRscores employer and a Compustat firm that the evidence could not settle.", f),
            ("conflict: the employer reaches two firms (both rows are listed); accept the right one (1) and reject the other (0).", f),
            ("single_source: only one source supports the link; accept (1) if it is the same company or a subsidiary it owned "
             "in 2012-2021, else reject (0).", f),
            ("Fill the yellow columns accept (1 or 0), decided_by and note. Leave accept empty when unsure.", f),
            ("Decide on names, sizes, company websites and 10-K subsidiary lists (Exhibit 21) only. Never use party, donation, "
             "DIPI or VRscores share data.", f),
            ("Then save this file in POLISY/linkage as hand_decisions.xlsx (a later round: hand_decisions_2.xlsx, and so on) "
             "and run build_links() again.", f)]
    for i, (s, font) in enumerate(text, start=1):
        hw.cell(row=i, column=1, value=s).font = font
    hw.column_dimensions["A"].width = 130
    ws = wb.create_sheet("Review")
    for j, (_, head, width) in enumerate(REVIEW_COLUMNS, start=1):
        cell = ws.cell(row=1, column=j, value=head)
        cell.font = fb
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[cell.column_letter].width = width
    for i, row in enumerate(r.itertuples(index=False), start=2):
        for j, (key, _, _) in enumerate(REVIEW_COLUMNS, start=1):
            v = getattr(row, key) if key in r.columns else None
            if isinstance(v, float) and math.isnan(v):
                v = None
            if key in ("vrid", "gvkey") and v is not None:
                v = str(v)
            if key == "vr_workers_max" and v is not None:
                v = int(v)
            if key == "pr_owner_share" and v is not None:
                v = round(float(v), 3)
            if key == "flag_patents_unowned":
                v = "yes" if v is True or str(v) == "True" else None
            cell = ws.cell(row=i, column=j, value=v)
            cell.font = f
            if key in ("vrid", "gvkey"):
                cell.number_format = "@"
            if key in ("accept", "decided_by", "note"):
                cell.fill = fill
    dv = DataValidation(type="list", formula1='"1,0"', allow_blank=True)
    ws.add_data_validation(dv)
    if len(r):
        dv.add(f"C2:C{len(r) + 1}")
    ws.freeze_panes = "F2"
    wb.save(path)
    log(f"review list: {len(r):,} pairs ({int(links.match_type.eq('conflict').sum()):,} in conflicts) in {path.name}. "
        "Decide them on names, sizes and company websites only; save the file as hand_decisions.xlsx")


def _sha256(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _freeze(out, links_path, S, note, inputs, norm_src):
    sha = _sha256(links_path)
    logf = out / "linkage_log.csv"
    hist = pd.read_csv(logf) if logf.exists() else pd.DataFrame(columns=["version", "created", "links_sha256", "note"])
    if len(hist) and hist.links_sha256.iloc[-1] == sha:
        version = int(hist.version.iloc[-1])
        log("links unchanged since the last freeze (same hash): the protocol version stays " + str(version))
    else:
        version = int(hist.version.max()) + 1 if len(hist) else 1
        row = {"version": version, "created": datetime.now(timezone.utc).isoformat(timespec="seconds"),
               "links_sha256": sha, "note": note}
        pd.concat([hist, pd.DataFrame([row])]).to_csv(logf, index=False)
    proto = {"version": version, "created": datetime.now(timezone.utc).isoformat(timespec="seconds"),
             "links_sha256": sha, "code": f"polisy_triangulate {__version__}", "name_cleaning": norm_src,
             "settings": {k: (list(v) if isinstance(v, tuple) else v) for k, v in S.items()}, "inputs": inputs,
             "rule": "No political variable was read while these links were built.", "note": note}
    (out / "protocol.json").write_text(json.dumps(proto, indent=2))
    return proto


def _frozen(out):
    """The frozen protocol, after checking that links_all.csv is the file that was frozen."""
    pf, lp = Path(out) / "protocol.json", Path(out) / "links_all.csv"
    if not pf.exists() or not lp.exists():
        raise RuntimeError("No frozen links yet: run build_links() first")
    proto = json.loads(pf.read_text())
    if _sha256(lp) != proto["links_sha256"]:
        raise RuntimeError("links_all.csv has changed since it was frozen. Do not edit it by hand: put decisions in "
                           "hand_decisions.xlsx and run build_links() again (it records a new version)")
    return proto


def _read_links(out):
    return pd.read_csv(Path(out) / "links_all.csv", dtype={"vrid": str, "gvkey": str, "pr_owner": str, "dn_gvkey": str},
                       low_memory=False)


def _readme(out):
    (Path(out) / "README_outputs.txt").write_text(
        """POLISY linkage: VRscores employers -> Compustat firms (GVKEY), triangulated and audited
Made by polisy_triangulate.py. Do not edit links_all.csv by hand: its hash is recorded in protocol.json.

links_all.csv          every candidate pair: evidence E1-E4, match type, the four link sets, and the details behind them
crosswalk_<set>.csv    the pairs of one link set (strict, main, broad, baseline), in module 04's format (accept = 1)
T1_match_types.csv     pairs, employers, firms and matched workers by match type (the paper's linkage table 1)
T1b_link_sets.csv      the same by link set, its share of all VRscores matched workers, and employers per firm
T3_agreement.csv       how often independent sources agree (the paper's linkage table 3)
review_list.xlsx       pairs needing a hand decision; fill accept (1 or 0), decided_by and note, save as hand_decisions.xlsx
hand_decisions*.xlsx   your decisions (you create them; a later round as hand_decisions_2.xlsx); build_links() reads them all
protocol.json          settings, input files, the links' hash and version: the frozen protocol
linkage_log.csv        every version of the links, with its hash and note
audit_key.csv          the audit sample with its match types (keep away from the coders)
audit_coder1.xlsx      the audit sheet for coder 1; audit_coder2.xlsx for coder 2 (identical, coded separately)
audit_resolution.csv   your agreed answers for the links the coders disagreed on (audit_id, final)
T4_audit_precision.csv precision by match type and by link set, with 95% intervals, and the coders' agreement
T2_selection.csv       linked versus unlinked firms on size, industry, patents and DIPI (made after the freeze)
workforce_<set>.csv/.dta  firm-year VRscores workforce for each link set (made after the freeze), for the sensitivity models
linkage_sensitivity.do Stata template that runs the thesis model once per link set
""")


# ----------------------------------------------------------------------------------------------- step B: the hand audit
AUDIT_COLUMNS = [("audit_id", "Audit ID", 10), ("vr_name", "VRscores employer", 42),
                 ("vr_workers_max", "VRscores matched workers (largest year)", 16), ("vr_years", "VRscores years", 12),
                 ("comp_name", "Compustat company", 40), ("comp_state", "Compustat state", 10),
                 ("comp_emp_match_year", "Compustat employees, 2015 (thousands)", 16),
                 ("correct", "Same company? 1 / 0 / ?", 14), ("evidence", "Where you checked", 30), ("note", "Note", 40)]


def draw_audit(base=None, paths=None, settings=None, n=None, overwrite=False):
    """Step B. Draw the blind audit sample (stratified by match type, seeded) and write two coder workbooks."""
    S = _settings(settings)
    P = _paths(base, paths)
    out = P["out"]
    proto = _frozen(out)
    key_path = out / "audit_key.csv"
    if key_path.exists() and not overwrite:
        log(f"the audit was already drawn ({key_path.name}); draw_audit(overwrite=True) draws a new one and makes "
            "the coding done so far useless")
        return pd.read_csv(key_path, dtype={"vrid": str, "gvkey": str})
    links = _read_links(out)
    univ = links[links.set_strict | links.set_main | links.set_broad | links.set_baseline].copy()
    univ = univ.sort_values(["match_type", "vrid", "gvkey"]).reset_index(drop=True)
    if univ.empty:
        log("audit: no links in any set yet, so there is nothing to audit (run build_links() first)")
        return pd.DataFrame()
    N = univ.match_type.value_counts().sort_index()
    n = int(min(n or S["audit_n"], len(univ)))
    alloc = {t: int(min(N[t], S["audit_min_per_type"])) for t in N.index}
    if sum(alloc.values()) > n:                                   # more types than the sample can cover at the minimum
        alloc = {t: max(1, int(n * N[t] / N.sum())) for t in N.index}
    rem = n - sum(alloc.values())
    while rem > 0:                                                # the rest in proportion to size (largest remainder)
        spare = {t: N[t] - alloc[t] for t in N.index if N[t] > alloc[t]}
        if not spare:
            break
        tot = sum(N[t] for t in spare)
        want = {t: rem * N[t] / tot for t in spare}
        add = {t: min(spare[t], int(math.floor(w))) for t, w in want.items()}
        if sum(add.values()) == 0:
            t = max(spare, key=lambda k: (want[k] - math.floor(want[k]), N[k], k))
            add = {t: 1}
        for t, a in add.items():
            alloc[t] += a
        rem = n - sum(alloc.values())
    rng = np.random.default_rng(S["seed"])
    parts = []
    for t in sorted(alloc):
        x = univ[univ.match_type == t]
        idx = rng.choice(len(x), size=alloc[t], replace=False)
        parts.append(x.iloc[np.sort(idx)].assign(stratum_links=int(N[t]), stratum_sampled=alloc[t]))
    smp = pd.concat(parts, ignore_index=True)
    smp = smp.iloc[rng.permutation(len(smp))].reset_index(drop=True)
    smp["audit_id"] = [f"A{i + 1:03d}" for i in range(len(smp))]
    smp["vr_years"] = smp.vr_first_year.astype("Int64").astype(str) + "-" + smp.vr_last_year.astype("Int64").astype(str)
    smp["comp_name"] = smp.comp_name_match_year.fillna(smp.comp_name)
    key = smp[["audit_id", "vrid", "gvkey", "match_type", "stratum_links", "stratum_sampled",
               "set_strict", "set_main", "set_broad", "set_baseline"]]
    key.to_csv(key_path, index=False)
    for coder in (1, 2):
        _audit_workbook(smp, out / f"audit_coder{coder}.xlsx", coder, proto)
    log(f"audit: {len(smp)} links drawn (" + ", ".join(f"{t} {alloc[t]}" for t in sorted(alloc)) + f"), seed {S['seed']}. "
        "Give audit_coder1.xlsx and audit_coder2.xlsx to two people who code them separately; keep audit_key.csv away "
        "from them. Then upload both back to POLISY/linkage and run audit_summary()")
    return key


def _audit_workbook(smp, path, coder, proto):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    hw = wb.active
    hw.title = "How to code"
    f = Font(name="Arial", size=10)
    fb = Font(name="Arial", size=10, bold=True)
    fill = PatternFill("solid", fgColor="FFF2CC")
    lines = [
        (f"Linkage audit, coder {coder}", Font(name="Arial", size=13, bold=True)),
        ("Prepared for the Research Team", f),
        (f"Links version {proto['version']} (hash {proto['links_sha256'][:12]}), drawn {proto['created'][:10]}", f),
        ("", f),
        ("What to do", fb),
        ("1. Open the sheet 'Audit'. Each row pairs a VRscores employer with a Compustat company.", f),
        ("2. Decide whether they are the same company. Fill only the yellow columns.", f),
        ("   1 = same company, or a division or subsidiary the Compustat company owned during 2012-2021", f),
        ("   0 = a different company (a namesake, a former parent, a firm sold before 2012)", f),
        ("   ? = you cannot tell after about five minutes of searching", f),
        ("3. Check the company's website, its LinkedIn page, Wikipedia, or the subsidiaries list in its 10-K "
         "(Exhibit 21, on SEC EDGAR). Write where you checked.", f),
        ("4. Work alone: do not discuss rows with the other coder until both of you have finished.", f),
        ("5. Judge identity only. Never look up or use anyone's politics (party, donations, DIPI, VRscores shares).", f),
        ("6. Save the file under the same name and send it back.", f),
        ("", f),
        ("Example (made up, to show the format)", fb),
        ("VRscores employer: 'Acme Robotics LLC' | Compustat company: 'ACME CORP' | correct: 1 | where: 10-K Exhibit 21 "
         "lists Acme Robotics LLC as a subsidiary", f),
    ]
    for i, (txt, font) in enumerate(lines, start=1):
        cell = hw.cell(row=i, column=1, value=txt)
        cell.font = font
    hw.cell(row=len(lines) + 2, column=1, value="Yellow cells = to fill in").fill = fill
    hw.column_dimensions["A"].width = 120

    ws = wb.create_sheet("Audit")
    for j, (_, head, width) in enumerate(AUDIT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=j, value=head)
        cell.font = fb
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[cell.column_letter].width = width
    for i, r in enumerate(smp.itertuples(index=False), start=2):
        for j, (key, _, _) in enumerate(AUDIT_COLUMNS, start=1):
            v = getattr(r, key) if key in smp.columns else None
            if isinstance(v, float) and math.isnan(v):
                v = None
            if key == "vr_workers_max" and v is not None:
                v = int(v)
            cell = ws.cell(row=i, column=j, value=v)
            cell.font = f
            if key in ("correct", "evidence", "note"):
                cell.fill = fill
    dv = DataValidation(type="list", formula1='"1,0,?"', allow_blank=True)
    ws.add_data_validation(dv)
    col = [k for k, _, _ in AUDIT_COLUMNS].index("correct") + 1
    letter = ws.cell(row=1, column=col).column_letter
    dv.add(f"{letter}2:{letter}{len(smp) + 1}")
    ws.freeze_panes = "B2"
    wb.save(path)


def _code(v):
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return np.nan
    s = str(v).strip().lower().replace(".0", "")
    return {"1": "1", "0": "0", "?": "?", "yes": "1", "no": "0", "y": "1", "n": "0"}.get(s, np.nan)


def _read_coder(path):
    d = pd.read_excel(path, sheet_name="Audit", dtype=str)
    d = d.rename(columns={head: key for key, head, _ in AUDIT_COLUMNS})
    if "audit_id" not in d.columns or "correct" not in d.columns:
        raise ValueError(f"{path.name}: the 'Audit' sheet has lost its columns; use the workbook as drawn")
    d["correct"] = d.correct.map(_code)
    return d[["audit_id", "correct"]]


def _wilson(k, n, z=1.96):
    if n == 0:
        return (np.nan, np.nan, np.nan)
    p = k / n
    den = 1 + z * z / n
    mid = (p + z * z / (2 * n)) / den
    half = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / den
    return p, max(0.0, mid - half), min(1.0, mid + half)


def _kappa(a, b):
    a, b = np.asarray(a), np.asarray(b)
    if len(a) == 0:
        return np.nan
    po = float(np.mean(a == b))
    pa, pb = float(np.mean(a == "1")), float(np.mean(b == "1"))
    pe = pa * pb + (1 - pa) * (1 - pb)
    return (po - pe) / (1 - pe) if pe < 1 else np.nan


def audit_summary(base=None, paths=None):
    """Step C. Inter-coder agreement, and precision by match type and by link set (stratified estimate, 95% interval)."""
    P = _paths(base, paths)
    out = P["out"]
    proto = _frozen(out)
    need = [out / "audit_key.csv", out / "audit_coder1.xlsx", out / "audit_coder2.xlsx"]
    gone = [f.name for f in need if not f.exists()]
    if gone:
        raise FileNotFoundError(f"missing in {out}: {', '.join(gone)}. Run draw_audit() first, then upload both coded "
                                "workbooks under their original names")
    key = pd.read_csv(out / "audit_key.csv", dtype={"vrid": str, "gvkey": str})
    c1, c2 = _read_coder(out / "audit_coder1.xlsx"), _read_coder(out / "audit_coder2.xlsx")
    m = key.merge(c1.rename(columns={"correct": "coder1"}), on="audit_id", how="left") \
           .merge(c2.rename(columns={"correct": "coder2"}), on="audit_id", how="left")
    done1, done2 = m.coder1.notna().sum(), m.coder2.notna().sum()
    both = m.coder1.isin(["1", "0"]) & m.coder2.isin(["1", "0"])
    agree = float((m.coder1[both] == m.coder2[both]).mean()) if both.any() else np.nan
    kappa = _kappa(m.coder1[both], m.coder2[both])
    m["final"] = np.where(m.coder1.notna() & (m.coder1 == m.coder2), m.coder1, np.nan)
    res = out / "audit_resolution.csv"
    if res.exists():
        r = pd.read_csv(res, dtype=str)
        r["final"] = r["final"].map(_code)
        m = m.merge(r[["audit_id", "final"]].rename(columns={"final": "resolved"}), on="audit_id", how="left")
        m["final"] = m.final.where(m.final.notna(), m.resolved)
    open_ = m[m.coder1.notna() & m.coder2.notna() & (m.coder1 != m.coder2) & m.final.isna()]
    open_[["audit_id", "coder1", "coder2"]].assign(final="").to_csv(out / "audit_disagreements.csv", index=False)
    log(f"audit: coder 1 coded {done1}/{len(m)}, coder 2 {done2}/{len(m)}; both coded 1 or 0 for {int(both.sum())}, "
        f"agreeing on {agree:.1%} (Cohen's kappa {kappa:.2f}); {len(open_)} disagreements open "
        "(audit_disagreements.csv: agree on an answer, save it as audit_resolution.csv with audit_id and final)")
    rows = []
    prec = {}
    for t, x in m.groupby("match_type"):
        coded = x[x.final.isin(["1", "0"])]
        k, n = int((coded.final == "1").sum()), len(coded)
        p, lo, hi = _wilson(k, n)
        prec[t] = (p, n, int(x.stratum_links.iloc[0]))
        rows.append({"level": "match type", "name": t, "links": int(x.stratum_links.iloc[0]), "sampled": len(x),
                     "coded": n, "unsure": int((x.final == "?").sum()), "correct": k, "precision": p,
                     "ci_low": lo, "ci_high": hi})
    links = _read_links(out)
    for s in SETS:
        x = links[links[f"set_{s}"]]
        Nh = x.match_type.value_counts()
        Ns = int(Nh.sum())
        est, var, cover, missing = 0.0, 0.0, 0, []
        for t, nt in Nh.items():
            if t not in prec or prec[t][1] == 0 or (isinstance(prec[t][0], float) and math.isnan(prec[t][0])):
                missing.append(t)
                continue
            p, n, Nt_all = prec[t]
            w = nt / Ns
            est += w * p
            fpc = max(0.0, 1 - n / Nt_all) if Nt_all else 1.0
            var += w * w * p * (1 - p) / n * fpc
            cover += nt
        if cover:
            est = est * Ns / cover                     # renormalized over the audited types
            se = math.sqrt(var) * Ns / cover
        else:
            est, se = np.nan, np.nan
        rows.append({"level": "link set", "name": s, "links": Ns, "sampled": np.nan, "coded": np.nan, "unsure": np.nan,
                     "correct": np.nan, "precision": est, "ci_low": max(0.0, est - 1.96 * se) if cover else np.nan,
                     "ci_high": min(1.0, est + 1.96 * se) if cover else np.nan,
                     "note": ("not audited: " + ", ".join(missing)) if missing else ""})
    rows.append({"level": "coders", "name": "agreement (both coded 1 or 0)", "links": int(both.sum()),
                 "precision": agree, "note": f"Cohen's kappa {kappa:.3f}"})
    T4 = pd.DataFrame(rows)
    T4.to_csv(out / "T4_audit_precision.csv", index=False)
    with pd.option_context("display.width", 160, "display.max_columns", 20):
        print(T4.round(3).to_string(index=False))
    log(f"precision table: T4_audit_precision.csv (links version {proto['version']})")
    return T4


# ----------------------------------------------------------------------------------------------- steps D-E: after the freeze
def _dipi(P):
    p = _resolve(P, "dipi")
    if p is not None:
        return pd.read_parquet(p)
    pc = _pc()
    if pc is not None:
        try:
            return pc.load_dipi()
        except Exception as e:  # noqa: BLE001
            log(f"DIPI: polisy_core could not load it ({type(e).__name__}: {e})")
    return None


def _col_like(cols, *tests):
    sq = {_squash(c): c for c in cols}
    for test in tests:
        hit = next((c for s, c in sq.items() if all(t in s for t in test)), None)
        if hit is not None:
            return hit
    return None


def selection_table(base=None, paths=None, settings=None):
    """Step D. Linked versus unlinked firms, for each link set: size, industry, patents and DIPI scores.
    Runs only on frozen links: DIPI is read here, never while links are built."""
    S = _settings(settings)
    P = _paths(base, paths)
    out = P["out"]
    proto = _frozen(out)
    links = _read_links(out)
    y0, y1 = S["window"]
    comp = _read_compustat(_compustat_source(P))
    cw = comp[(comp.fyear >= y0) & (comp.fyear <= y1)]
    f = cw.groupby("gvkey").agg(emp=("emp", "mean"), sale=("sale", "mean"), assets=("at", "mean"),
                                sich=("sich", "first"))
    f["log_employees"] = np.log(f["emp"].where(f["emp"] > 0) * 1000)
    f["log_sales"] = np.log(f["sale"].where(f["sale"] > 0))
    f["log_assets"] = np.log(f["assets"].where(f["assets"] > 0))
    f["manufacturing"] = np.where(f["sich"].notna(), f["sich"].between(2000, 3999).astype(float), np.nan)
    variables = ["log_employees", "log_sales", "log_assets", "manufacturing"]
    fp = _resolve(P, "firm_patents")
    if fp is not None:
        pt = pd.read_parquet(fp)
        pcol = "pat_filed" if "pat_filed" in pt.columns else _col_like(pt.columns, ("pat",))
        pt["gvkey"] = _gvkey(pt.gvkey)
        pt = pt[(pd.to_numeric(pt.year, errors="coerce") >= y0) & (pd.to_numeric(pt.year, errors="coerce") <= y1)]
        f["patents_per_year"] = pt.groupby("gvkey")[pcol].sum() / (y1 - y0 + 1)
        f["patents_per_year"] = f.patents_per_year.fillna(0)
        f["log1p_patents"] = np.log1p(f.patents_per_year)
        variables += ["log1p_patents"]
    d = _dipi(P)
    universe = "Compustat firms in the window"
    if d is not None:
        d = d.copy()
        d["gvkey"] = _gvkey(d.gvkey)
        d = d[(pd.to_numeric(d.year, errors="coerce") >= y0) & (pd.to_numeric(d.year, errors="coerce") <= y1)]
        emp_col = _col_like(d.columns, ("emp", "liberal", "10"), ("emp", "liberal"))
        ceo_col = _col_like(d.columns, ("ceo", "liberal", "10"), ("ceo", "liberal"))
        for name, col in (("dipi_employee_liberalism", emp_col), ("dipi_ceo_liberalism", ceo_col)):
            if col is not None:
                f[name] = pd.to_numeric(d[col], errors="coerce").groupby(d.gvkey).mean()
                variables.append(name)
        f = f[f.index.isin(set(d.gvkey))]
        universe = "DIPI firms (the thesis universe)"
    rows = []
    for s in SETS:
        linked = f.index.isin(set(links.loc[links[f"set_{s}"], "gvkey"]))
        for v in variables:
            a, b = f.loc[linked, v].dropna(), f.loc[~linked, v].dropna()
            sd = math.sqrt((a.var() + b.var()) / 2) if len(a) > 1 and len(b) > 1 else np.nan
            rows.append({"set": s, "variable": v, "linked_mean": a.mean(), "linked_sd": a.std(), "linked_n": len(a),
                         "unlinked_mean": b.mean(), "unlinked_sd": b.std(), "unlinked_n": len(b),
                         "std_difference": (a.mean() - b.mean()) / sd if sd else np.nan})
    T2 = pd.DataFrame(rows)
    T2.to_csv(out / "T2_selection.csv", index=False)
    log(f"selection table: T2_selection.csv ({universe}, {y0}-{y1}; links version {proto['version']}). "
        "Standardized differences above 0.25 in absolute value deserve a sentence in the paper")
    return T2


def workforce_panels(base=None, paths=None, write_do=True):
    """Step E. The firm-year VRscores workforce for each link set (module 06's definitions), as CSV and Stata files.
    Reads VRscores party counts, so it runs only on frozen links."""
    P = _paths(base, paths)
    out = P["out"]
    proto = _frozen(out)
    links = _read_links(out)
    vr_path = _resolve(P, "vr_employer", True)
    c = _con(out)
    try:
        cols = _cols(c, vr_path)
        m = {"vrid": _pick(cols, "vrid", "unit"), "year": _pick(cols, "year"), "workers": _pick(cols, "employeecount", "workers"),
             "dem": _pick(cols, "demworkersimp", "dem"), "rep": _pick(cols, "repworkersimp", "rep"),
             "dem_raw": _pick(cols, "demworkersraw", "demraw"), "rep_raw": _pick(cols, "repworkersraw", "repraw")}
        if any(m[k] is None for k in ("vrid", "year", "workers", "dem", "rep")):
            raise ValueError(f"VRscores employer panel lacks party counts; its columns are {cols[:20]}")
        q = lambda x: '"' + x.replace('"', '""') + '"'  # noqa: E731
        raw = (f", sum(CAST(e.{q(m['dem_raw'])} AS DOUBLE)) AS dem_raw, sum(CAST(e.{q(m['rep_raw'])} AS DOUBLE)) AS rep_raw"
               if m["dem_raw"] and m["rep_raw"] else ", NULL::DOUBLE AS dem_raw, NULL::DOUBLE AS rep_raw")
        for s in SETS:
            x = links.loc[links[f"set_{s}"], ["vrid", "gvkey", "match_type"]]
            c.register("set_df", x)
            panel = c.execute(f"""SELECT l.gvkey, CAST(e.{q(m['year'])} AS INTEGER) AS year, count(*) AS n_employers,
                                         count(*) FILTER (WHERE l.match_type = 'triangulated_patents') AS n_triangulated,
                                         sum(CAST(e.{q(m['workers'])} AS DOUBLE)) AS workers,
                                         sum(CAST(e.{q(m['dem'])} AS DOUBLE)) AS dem, sum(CAST(e.{q(m['rep'])} AS DOUBLE)) AS rep
                                         {raw}
                                  FROM read_parquet('{_sql(vr_path)}') e JOIN set_df l
                                    ON CAST(e.{q(m['vrid'])} AS VARCHAR) = l.vrid
                                  GROUP BY 1, 2 ORDER BY 1, 2""").df()
            two = panel.dem + panel.rep
            panel["vr_dem_share"] = panel.dem / two.where(two > 0)
            panel["vr_rep_share"] = panel.rep / two.where(two > 0)
            two_raw = panel.dem_raw + panel.rep_raw
            panel["vr_dem_share_raw"] = panel.dem_raw / two_raw.where(two_raw > 0)
            panel["gvkey_n"] = pd.to_numeric(panel.gvkey, errors="coerce").astype("Int64")
            panel["link_set"] = s
            panel["links_version"] = int(proto["version"])
            panel.to_csv(out / f"workforce_{s}.csv", index=False)
            st = panel.copy()
            st["gvkey_n"] = st.gvkey_n.astype("float")
            st.to_stata(out / f"workforce_{s}.dta", write_index=False, version=118,
                        variable_labels={"vr_dem_share": "VRscores Democratic share of two-party workers (imputed)",
                                         "n_triangulated": "employers linked by Compustat name and patent route"})
            log(f"workforce panel {s}: {panel.gvkey.nunique():,} firms, {len(panel):,} firm-years -> workforce_{s}.csv/.dta")
    finally:
        c.close()
    if write_do:
        write_stata_template(out)
    return True


STATA_TEMPLATE = r"""* linkage_sensitivity.do: the thesis model once per link set (made by polisy_triangulate.py)
* Before running: edit the five lines marked EDIT. Run it only after the links are frozen and the audit is coded.
clear all
set more off

global LINK   "C:/path/to/POLISY/linkage"            // EDIT: the linkage folder (download it from Google Drive)
global THESIS "C:/path/to/thesis_firm_year.dta"      // EDIT: your firm-year panel with gvkey, year, patents, DIPI, controls
global CEO    ceo_liberalism                         // EDIT: your DIPI CEO liberalism variable (0 = Republican, 1 = Democratic)
global controls "lnsize leverage capint firmage psm"  // EDIT: your control variables
global MODEL  "nbreg patents incongruence $controls i.year, vce(cluster gvkey)"   // EDIT: Table 5, Model 1 of the thesis

* Your thesis file needs one row per gvkey and year (variables named gvkey and year). It may store gvkey as a number or
* as text: the workforce files carry both (gvkey as six-digit text, gvkey_n as a number).
foreach set in strict main broad baseline {
    use "$THESIS", clear
    capture confirm numeric variable gvkey
    if !_rc {
        rename gvkey gvkey_n
        merge 1:1 gvkey_n year using "$LINK/workforce_`set'.dta", keep(match) nogenerate
    }
    else {
        replace gvkey = string(real(gvkey), "%06.0f") if real(gvkey) < .   // "1004" -> "001004", as in Compustat
        merge 1:1 gvkey year using "$LINK/workforce_`set'.dta", keep(match) nogenerate
    }
    capture drop incongruence
    generate incongruence = abs($CEO - vr_dem_share)
    display as text _n "==== link set: `set' ===="
    $MODEL
    estimates store m_`set'
}
estimates table m_strict m_main m_broad m_baseline, keep(incongruence) b(%9.3f) se(%9.3f) stats(N)
* The coefficient on incongruence should be similar across the four sets. Report all four in the linkage appendix.
"""


def write_stata_template(out=None, base=None):
    out = Path(out or _paths(base)["out"])
    p = out / "linkage_sensitivity.do"
    p.write_text(STATA_TEMPLATE)
    log(f"Stata template: {p.name} (edit the lines marked EDIT)")
    return p


# ----------------------------------------------------------------------------------------------- command line
if __name__ == "__main__":
    import argparse

    ap = argparse.ArgumentParser(description="Triangulated VRscores -> Compustat linkage")
    ap.add_argument("step", choices=["links", "audit", "summary", "selection", "panels"])
    ap.add_argument("--base", default=None, help="your POLISY folder (default /content/drive/MyDrive/POLISY)")
    ap.add_argument("--note", default="", help="why this version of the links was built (links step)")
    a = ap.parse_args()
    {"links": lambda: build_links(a.base, note=a.note), "audit": lambda: draw_audit(a.base),
     "summary": lambda: audit_summary(a.base), "selection": lambda: selection_table(a.base),
     "panels": lambda: workforce_panels(a.base)}[a.step]()
